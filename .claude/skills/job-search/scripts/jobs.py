"""Find, merge, check and rank jobs for Veer, then write the Excel workbook and job board.

usage: python jobs.py fetch [--days 7] [--sources ba,freehire,arbeitnow] --out jobs/2026-10-26
       python jobs.py pool DUMP.json [DUMP.json ...] --out jobs/2026-10-26 [--top 150]
       python jobs.py details ID [ID ...] --out jobs/2026-10-26
       python jobs.py export --out jobs/2026-10-26
       python jobs.py --selftest

fetch    searches the free sources (Bundesagentur für Arbeit, freehire, Arbeitnow) and writes
         OUT/free_sources.json in the Apify actor's item format, so pool takes it like a LinkedIn dump.
pool     merges the dumps (the Apify ones get-dataset-items saves to files, plus free_sources.json), drops
         duplicates, recruiting agencies, senior roles and jobs he already applied to, saves OUT/pool.json and
         prints the best candidates.
details  checks each job on its source (still open? company apply link?) and prints the posting lines that
         scoring needs: tasks, requirements, language, experience, hours, start and salary.
export   ranks OUT/picks.json (the agent's scores) and writes OUT/jobs.json (read by /job-apply),
         OUT/Veer_Job_Search_<date>.xlsx and OUT/job_board.html.
"""
import argparse
import base64
import datetime as dt
import html
import json
import os
import re
import shutil
import ssl
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

SKILL = Path(__file__).resolve().parents[1]
PROJECT = Path(__file__).resolve().parents[4]
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                         "Chrome/126.0 Safari/537.36", "Accept-Language": "en-US,en;q=0.9,de;q=0.8"}
WEIGHTS = {"tech": 30, "experience": 25, "education": 15, "career": 15, "eligibility": 10, "freshness": 5}

# first match wins, so a "Werkstudent Masterarbeit" counts as a thesis
CATEGORIES = [
    ("Thesis", r"masterarbeit|master'?s? ?thesis|abschlussarbeit|thesis|masterand"),
    ("Research", r"hilfskraft|hiwi|\bshk\b|\bwhk\b|research (?:assistant|intern|student)|student research|"
                 r"wissenschaftliche|forschungsassist"),
    ("Werkstudent", r"werkstudent|working student|student employee|studentische|student assistant|"
                    r"student (?:developer|engineer|job)|studentenjob"),
    ("Internship", r"praktik|\bintern\b|internship|praxissemester"),
    ("Trainee", r"trainee|graduate program|nachwuchs|einstiegsprogramm|direkteinstieg"),
]
DOMAINS = [
    ("AI/ML", r"\bki\b|\bai\b|künstliche intelligenz|machine learning|\bml\b|deep learning|\bllm|genai|generative|"
              r"\bnlp\b|computer vision|bildverarbeitung|data scien"),
    ("Data", r"\bdata\b|daten|analytics|analyst|business intelligence|\bbi\b|power bi|reporting|\betl\b|\bsql\b"),
    ("Cloud/DevOps", r"cloud|devops|platform engineer|\bsre\b|kubernetes|infrastru|mlops|\baws\b|azure"),
    ("Security", r"security|sicherheit|\bsoc\b|cyber|\biam\b"),
    ("IoT/Embedded/Robotics", r"embedded|\biot\b|robot|firmware|automatisierung|automation|sensor|steuerung"),
    ("QA/Testing", r"\bqa\b|\btest|quality engineer|qualitätssicherung|validation|verifikation"),
    ("Consulting/Enterprise IT", r"consult|berater|\bsap\b|\berp\b|\bcrm\b|business analyst|digitali|transformation"),
    ("Research", r"research|forschung|wissenschaft"),
    ("Software", r"software|develop|entwickl|engineer|programm|full ?stack|backend|frontend|\bweb|\bapp\b|python|java"),
]
TECH_TITLE = re.compile(
    r"software|develop|entwickl|engineer|ingenieur|informatik|\bit\b|\bit-|data|daten|\bki\b|\bai\b|machine learning|"
    r"\bml\b|cloud|devops|cyber|security|sicherheit|\bsap\b|\bbi\b|business intelligence|analytics|analyst|embedded|"
    r"\biot\b|robot|automatisierung|automation|digital|\btest|\bqa\b|\bweb|frontend|backend|full ?stack|python|java|"
    r"\.net|c\+\+|research|forschung|thesis|masterarbeit|abschlussarbeit|hilfskraft|hiwi|platform|infrastru|system|"
    r"network|netzwerk|database|datenbank|\bllm|genai|computer|tech|scien")
JUNIOR = re.compile(r"junior|graduate|entry|einsteiger|absolvent|trainee|berufseinstieg|young professional")
AGENCY = re.compile(r"recruit|personaldienst|personalvermittlung|zeitarbeit|staffing|headhunt|\bhays\b|randstad|"
                    r"adecco|manpower|\bgulp\b|amadeus fire|experis|orizon|jobvalley|studitemps|zenjob|jobster|"
                    r"jobbird|academic work")
GENDER = re.compile(r"\((?:[mwfdx]\s*/\s*){1,3}[mwfdx]\)|\(all genders?\)|\(gn\)|\(d\)|\*in\b|:in\b|/in\b|\*")
STOP = set("werkstudent werkstudentin working student intern internship praktikum junior m w d f gn all genders "
           "und and im in for of the der die das mit with as als".split())
# company-name words that also occur in role names, so they can't tell companies apart
GENERIC = set("gmbh ag se kg co group gruppe germany deutschland deutsche deutscher verlag international technologies "
              "technology tech solutions services service systems holding digital consulting energy europe data ai it "
              "ki software analytics cloud labs engineering management technical technische university universität "
              "institute handels product development sales operations compliance governance automation platform".split())

# section headings, matched at the start of a short line or of a line ending in ":"
REQ_HEAD = re.compile(
    r"(?:dein |ihr |your )?(?:profil|profile|qualifikation|qualifications?|requirements?|skills|background)|"
    r"(?:das|was) (?:bringst du mit|du mitbringst|sie mitbringen|bringen sie mit)|what you (?:bring|need|have|offer)|"
    r"we(?:'re| are) looking (?:for|forward to)|who you are|about you|über dich|(?:was )?wir erwarten|voraussetzungen|"
    r"anforderungen|must[- ]haves?|you have|you bring|das zeichnet dich aus|was dich auszeichnet|this is you|"
    r"what we expect|wen wir suchen", re.I)
TASK_HEAD = re.compile(
    r"(?:deine |ihre |your |key |main )?(?:aufgaben|tasks|mission|role|responsibilities|tätigkeiten|impact)|"
    r"what you(?:'ll| will) (?:do|be doing)|you can look forward to these tasks|(?:das|was) (?:dich|sie) erwartet|"
    r"dein job|the role", re.I)
OTHER_HEAD = re.compile(
    r"(?:was |das )?wir bieten|das bieten wir|unser angebot|our offer|we offer|benefits|what we offer|why (?:us|join)|"
    r"über uns|about us|kontakt|contact|deine vorteile|ihre vorteile|perks|interesse|bewirb dich|apply now|"
    r"your working environment|what you can expect|weitere informationen|additional information|einleitung|"
    r"introduction", re.I)
CLOSED_PAGE = re.compile(r"no longer (?:available|accepting|active)|position (?:has been|is) (?:filled|closed)|"
                         r"nicht mehr (?:verfügbar|aktiv|ausgeschrieben)|bereits besetzt|"
                         r"stelle ist (?:leider )?(?:nicht mehr|bereits)", re.I)
DE_WORDS = re.compile(r"\b(?:und|der|die|das|mit|für|wir|sie|du|dich|ist|nicht|auf|bei|oder|unser|deine|ihre)\b", re.I)
EN_WORDS = re.compile(r"\b(?:and|the|with|for|we|you|is|not|on|at|or|our|your|are|will)\b", re.I)
LEGAL = set("gmbh ag se kg co mbh kgaa ug inc ltd llc group gruppe holding deutschland germany".split())

BA_API = "https://rest.arbeitsagentur.de/jobboerse/jobsuche-service"
BA_KEY = {"X-API-Key": "jobboerse-jobsuche"}  # the public client ID from the API's own documentation
FREEHIRE = "https://freehire.me/api/v1"
ARBEITNOW = "https://www.arbeitnow.com/api/job-board-api"
SOURCES = {"ba": "Bundesagentur für Arbeit", "freehire": "freehire", "arbeitnow": "Arbeitnow",
           "indeed": "Indeed", "stepstone": "StepStone", "linkedin": "LinkedIn"}
FREE_SOURCES = ["ba", "freehire", "arbeitnow"]
# Apify actors (pay per result); these sources need an Apify API token
APIFY = "https://api.apify.com/v2"
ACTORS = {"linkedin": "cheap_scraper/linkedin-job-scraper", "indeed": "valig/indeed-jobs-scraper",
          "stepstone": "valig/stepstone-jobs-scraper"}
PRICE = {"linkedin": 0.0007, "indeed": 0.0001, "stepstone": 0.0004}  # USD per job on Apify's free tier
LINKEDIN_FIELDS = ("jobId,jobTitle,companyName,location,publishedAt,postedTime,experienceLevel,contractType,"
                   "applicationsCount,applyType,keywordMatchScorePercentage,dynamicFilterMatch,salaryInfo,sector,"
                   "searchString,yearsOfExperience,jobDescription")
# Indeed and StepStone take one search per actor run, so they get a short list of broad searches
BOARD_QUERIES = ["Werkstudent Informatik", "Werkstudent Data", "Werkstudent Software", "Werkstudent KI",
                 "Working Student AI", "Praktikum Informatik", "Masterarbeit Informatik", "Junior Data",
                 "Junior Software"]
STATE_CODES = {"BW": "Baden-Württemberg", "BY": "Bavaria", "BE": "Berlin", "BB": "Brandenburg", "HB": "Bremen",
               "HH": "Hamburg", "HE": "Hesse", "MV": "Mecklenburg-West Pomerania", "NI": "Lower Saxony",
               "NW": "North Rhine-Westphalia", "RP": "Rhineland-Palatinate", "SL": "Saarland", "SN": "Saxony",
               "ST": "Saxony-Anhalt", "SH": "Schleswig-Holstein", "TH": "Thuringia"}
# freehire searches titles for these; one OR group may hold at most five alternatives
FREEHIRE_QUERIES = [
    'title:werkstudent OR title:werkstudentin OR title:"working student" OR title:"student assistant" OR title:studentische',
    "title:praktikum OR title:praktikant OR title:intern OR title:internship OR title:trainee",
    "title:masterarbeit OR title:thesis OR title:abschlussarbeit OR title:hilfskraft OR title:hiwi",
    'title:junior OR title:graduate OR title:absolvent OR title:berufseinsteiger OR title:"entry level"',
]
SENIOR_LEVELS = {"senior", "lead", "staff", "principal", "c_level"}
STATES = {"badenwuerttemberg": "Baden-Württemberg", "bayern": "Bavaria", "berlin": "Berlin",
          "brandenburg": "Brandenburg", "bremen": "Bremen", "hamburg": "Hamburg", "hessen": "Hesse",
          "mecklenburgvorpommern": "Mecklenburg-West Pomerania", "niedersachsen": "Lower Saxony",
          "nordrheinwestfalen": "North Rhine-Westphalia", "rheinlandpfalz": "Rhineland-Palatinate",
          "saarland": "Saarland", "sachsen": "Saxony", "sachsenanhalt": "Saxony-Anhalt",
          "schleswigholstein": "Schleswig-Holstein", "thueringen": "Thuringia"}
# freehire gives a city without its state; these are the cities most job ads name
CITY_STATES = {
    **dict.fromkeys(["munich", "münchen", "nuremberg", "nürnberg", "augsburg", "regensburg", "ingolstadt", "würzburg",
                     "erlangen", "garching", "unterföhring", "fürth", "bamberg", "passau", "landshut"], "Bavaria"),
    **dict.fromkeys(["stuttgart", "karlsruhe", "mannheim", "heidelberg", "freiburg", "ulm", "heilbronn", "böblingen",
                     "sindelfingen", "walldorf", "tübingen", "reutlingen", "ditzingen", "pforzheim"], "Baden-Württemberg"),
    **dict.fromkeys(["cologne", "köln", "düsseldorf", "dortmund", "essen", "duisburg", "bochum", "bonn", "aachen",
                     "münster", "bielefeld", "wuppertal", "paderborn", "siegen", "gelsenkirchen", "mönchengladbach",
                     "krefeld", "neuss"], "North Rhine-Westphalia"),
    **dict.fromkeys(["frankfurt", "frankfurt am main", "wiesbaden", "darmstadt", "offenbach", "kassel", "eschborn",
                     "bad homburg"], "Hesse"),
    **dict.fromkeys(["hanover", "hannover", "brunswick", "braunschweig", "wolfsburg", "göttingen", "oldenburg",
                     "osnabrück", "lüneburg"], "Lower Saxony"),
    **dict.fromkeys(["leipzig", "dresden", "chemnitz"], "Saxony"),
    **dict.fromkeys(["magdeburg", "halle", "halle (saale)"], "Saxony-Anhalt"),
    **dict.fromkeys(["erfurt", "jena", "weimar"], "Thuringia"),
    **dict.fromkeys(["potsdam", "cottbus"], "Brandenburg"),
    **dict.fromkeys(["kiel", "lübeck", "flensburg"], "Schleswig-Holstein"),
    **dict.fromkeys(["rostock", "schwerin"], "Mecklenburg-West Pomerania"),
    **dict.fromkeys(["mainz", "koblenz", "trier", "kaiserslautern", "ludwigshafen"], "Rhineland-Palatinate"),
    **dict.fromkeys(["saarbrücken"], "Saarland"),
    **dict.fromkeys(["bremerhaven"], "Bremen"),
}
JOB_BOARDS = re.compile(r"arbeitnow\.|linkedin\.|indeed\.|stepstone\.|arbeitsagentur\.|xing\.|glassdoor\.", re.I)
# Windows' own curl first: it fetches the intermediate certificate rest.arbeitsagentur.de leaves out
CURL = shutil.which("curl", path=r"C:\Windows\System32") or shutil.which("curl")


def simple(s: str) -> str:
    return re.sub(r"[^a-z0-9äöüß+#]+", " ", s.lower()).strip()


def category(title: str) -> str:
    t = title.lower()
    return next((name for name, rx in CATEGORIES if re.search(rx, t)), "Full-time")


def domain(title: str, desc: str) -> str:
    for text in (title.lower(), desc[:1500].lower()):
        hit = next((name for name, rx in DOMAINS if re.search(rx, text)), None)
        if hit:
            return hit
    return "Other"


def posting_language(text: str) -> str:
    de, en = len(DE_WORDS.findall(text)), len(EN_WORDS.findall(text))
    return "German" if de > en else "English" if en else "Unknown"


def min_years(values: list):
    """Smallest year count from strings like "2-5", "3+" or "mindestens 3"; None if there are none."""
    nums = [int(m.group()) for v in values for m in [re.search(r"\d+", str(v))] if m]
    nums = [n for n in nums if n <= 15]  # "100 Jahre Firmengeschichte" is not an experience requirement
    return min(nums) if nums else None


def years_in_text(desc: str) -> list:
    pats = [r"(\d+)\s*\+?\s*(?:-|–|bis|to)?\s*\d*\s*\+?\s*(?:years?|jahre?n?)\b[^.\n]{0,60}?(?:experience|erfahrung)",
            r"(?:experience|erfahrung)[^.\n]{0,40}?(\d+)\s*\+?\s*(?:years?|jahre?n?)\b"]
    return [m.group(1) for p in pats for m in re.finditer(p, desc, re.I)]


def age_from_text(posted: str) -> int:
    m = re.search(r"(\d+)\s*(minute|hour|day|week|month)", posted)
    if not m:
        return 99
    n, unit = int(m.group(1)), m.group(2)
    return {"minute": 0, "hour": 0, "day": n, "week": 7 * n, "month": 30 * n}[unit]


def place(loc: str) -> str:
    """"Munich, Bavaria, Germany" -> "Munich, Bavaria"; "Berlin, Berlin, Germany" -> "Berlin"."""
    parts = [p.strip() for p in loc.split(",") if p.strip() and p.strip() != "Germany"]
    if len(parts) == 2 and parts[0] == parts[1]:
        parts = parts[:1]
    return ", ".join(parts) or "Germany (nationwide)"


def state(loc: str) -> str:
    parts = [p.strip() for p in loc.split(",")]
    return parts[-2] if len(parts) >= 3 else ""


def state_label(loc: str) -> str:
    return state(loc) or ("nationwide/remote" if place(loc).startswith("Germany") else "state not given")


def english_state(name: str) -> str:
    """A German state as LinkedIn writes it: BA's "NORDRHEIN_WESTFALEN" or Arbeitnow's "Bayern" -> English."""
    key = re.sub(r"[^a-z]", "", name.lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue"))
    return STATES.get(key, "")


def german_location(city: str, region: str = "") -> str:
    """LinkedIn's "City, State, Germany" form, so place() and state() read every source alike."""
    city = city.split(",")[0].strip()  # the Bundesagentur sometimes writes "Kaufungen, Hessen"
    st = english_state(region) or english_state(city) or CITY_STATES.get(city.lower(), "")
    return ", ".join(p for p in (city, st, "Germany") if p)


def company_key(name: str) -> str:
    """"Scalable Capital", "ScalableGmbH" and "scalable-gmbh" all start "scalable": freehire often names
    companies by their career-site slug."""
    key = "".join(w for w in simple(name).split() if w not in LEGAL)
    return re.sub(r"(?:gmbh|mbh|ag|se|kg|ug|inc|ltd|llc)+$", "", key) or key


def same_company(a: str, b: str) -> bool:
    n = min(len(a), len(b), 6)
    return n >= 3 and a[:n] == b[:n]


def to_job(x: dict, today: dt.date) -> dict:
    desc = x.get("jobDescription") or ""
    title = (x.get("jobTitle") or "").strip()
    published = (x.get("publishedAt") or "")[:10]
    age = (today - dt.date.fromisoformat(published)).days if published else age_from_text(x.get("postedTime") or "")
    years = [y.get("years", "") for y in x.get("yearsOfExperience") or []] + years_in_text(desc)
    src = x.get("source") or "LinkedIn"
    linkedin = f"https://www.linkedin.com/jobs/view/{x['jobId']}/" if src == "LinkedIn" else ""
    return {"id": str(x["jobId"]), "source": src, "title": title, "company": (x.get("companyName") or "").strip(),
            "locations": [x.get("location") or "Germany"], "published": published, "age": age,
            "posted_at": x.get("publishedAt") or "",  # with the time of day where the source gives one
            "level": x.get("experienceLevel") or "", "contract": x.get("contractType") or "",
            "applicants": x.get("applicationsCount") or "", "apply_type": x.get("applyType") or "",
            "linkedin": linkedin, "url": linkedin or x.get("jobUrl") or "", "company_url": x.get("company_url"),
            "ref": x.get("ref"), "slug": x.get("slug"), "kw": round(x.get("keywordMatchScorePercentage") or 0),
            "agency_flag": x.get("dynamicFilterMatch") is False, "salary_info": x.get("salaryInfo") or [],
            "sector": x.get("sector") or "", "search": x.get("searchString") or "", "category": category(title),
            "domain": domain(title, desc), "lang": posting_language(desc), "min_years": min_years(years),
            "description": desc}


def dedupe(jobs: list) -> list:
    """One entry per vacancy: same ID, or same company and title on any source (then the locations are merged).
    The LinkedIn copy stays the main record, and a company apply link found on another source is kept."""
    by_id = {}
    for j in jobs:
        by_id.setdefault(j["id"], j)
    by_title = {}
    for j in sorted(by_id.values(), key=lambda j: (j.get("source", "LinkedIn") != "LinkedIn", j["age"])):
        same_title = by_title.setdefault(simple(GENDER.sub(" ", j["title"].lower())), [])
        m = next((m for m in same_title if same_company(company_key(m["company"]), company_key(j["company"]))), None)
        if m is None:
            same_title.append(j)
            continue
        city = lambda loc: place(loc).split(",")[0].lower()[:8]  # "Bad Homburg" and "Bad Homburg vor der Höhe"
        m["locations"] += [l for l in j["locations"] if all(city(l) != city(e) for e in m["locations"])]
        m["company_url"] = m.get("company_url") or j.get("company_url")
        if j.get("source", "LinkedIn") != m.get("source", "LinkedIn") and j["source"] not in m.setdefault("also_on", []):
            m["also_on"].append(j["source"])
    return [j for group in by_title.values() for j in group]


def applied_sets() -> list:
    """Word sets from his tailored CV file names, e.g. CV_DEAG_Werkstudent_KI_Automatisierung, including the
    ones the desktop app writes into applications/."""
    return [set(p.stem.lower().split("_")[4:]) - STOP for p in PROJECT.rglob("Veer_Pratap_Singh_CV_*.tex")]


def applied_match(job: dict, sets: list):
    """"job" if the company and two title words match one of his tailored CVs, "company" if only the company does.
    A German repost of a job he applied to in English only matches "company", so the agent checks those by eye."""
    company = set(simple(job["company"]).split()) - GENERIC - STOP
    title = set(simple(GENDER.sub(" ", job["title"].lower())).split()) - STOP
    hits = [s for s in sets if company & s]
    if any(len(title & s) >= 2 for s in hits):
        return "job"
    return "company" if hits else None


def drop_reason(j: dict, applied: list):
    t, c = j["title"].lower(), j["company"].lower()
    if "(inaktiv)" in c:
        return "inactive company page"
    if j["agency_flag"] or AGENCY.search(c):
        return "recruiting agency"
    if j["level"] in ("Director", "Executive"):
        return "senior level"
    if j["min_years"] is not None and j["min_years"] >= 3:
        return "3+ years of experience"
    if re.search(r"bachelorarbeit|bachelor'?s? thesis", t) and "master" not in t:
        return "bachelor thesis"
    if re.search(r"\bph\.?\s?d\b|doktorand|postdoc|doctoral|promotionsstelle", t):
        return "PhD or postdoc"
    if re.search(r"fachinformatiker|\(b\.?\s?(?:sc|a|eng)\.?\)", t):
        return "apprenticeship or bachelor programme"
    if not TECH_TITLE.search(t) and j["kw"] < 40:
        return "not a tech role"
    if applied_match(j, applied) == "job":
        return "already applied"
    return None


def prescore(j: dict) -> int:
    """Rough order for the review list only; the agent scores the shortlist properly."""
    t = j["title"].lower()
    s = j["kw"] * 0.5 + (25 if TECH_TITLE.search(t) else 0)
    s += 15 if j["age"] <= 2 else 10 if j["age"] <= 4 else 5 if j["age"] <= 7 else 0
    s += 10 if j["category"] != "Full-time" or JUNIOR.search(t) else 0
    s -= 15 if j["level"] == "Mid-Senior level" else 0
    s -= 10 if (j["min_years"] or 0) >= 2 else 0
    return round(s)


def load_items(paths: list) -> list:
    items = []
    for p in paths:
        data = json.loads(Path(p).read_text(encoding="utf-8"))
        items += data["items"] if isinstance(data, dict) else data
    return items


def make_pool(raw: list, today: dt.date, prev: dict = None) -> tuple:
    """(pool, skipped_applied): the merged, filtered and pre-scored jobs. prev is an earlier pool whose
    checks (status, description, snippets) are kept, so a second search pass doesn't redo details."""
    jobs = dedupe([to_job(x, today) for x in raw if x.get("jobId")])
    applied, dropped, skipped_applied, kept = applied_sets(), Counter(), [], []
    for j in jobs:
        reason = drop_reason(j, applied)
        if reason == "already applied":
            skipped_applied.append(f"{j['title']} | {j['company']}")
        if reason:
            dropped[reason] += 1
            continue
        j["prescore"], j["applied_company"] = prescore(j), applied_match(j, applied) == "company"
        kept.append(j)
    kept.sort(key=lambda j: -j["prescore"])
    meta = {"date": today.isoformat(), "retrieved": len(raw), "unique": len(jobs),
            "retrieved_by_source": dict(Counter(x.get("source") or "LinkedIn" for x in raw)),
            "kept_by_source": dict(Counter(j["source"] for j in kept)),
            "searches_with_results": len({x.get("searchString") for x in raw if not x.get("source")}),
            "dropped": dict(dropped), "kept": len(kept), "categories": dict(Counter(j["category"] for j in kept)),
            "domains": dict(Counter(j["domain"] for j in kept)),
            "states": dict(Counter(state_label(l) for j in kept for l in j["locations"])),
            "languages": dict(Counter(j["lang"] for j in kept)), "assessed": []}
    if prev:
        ids = {j["id"] for j in kept}
        for j in kept:
            p = prev["jobs"].get(j["id"], {})
            j.update({k: p[k] for k in ("status", "apply_url", "apply_check", "checked", "snippets",
                                        "workplace_hint") if k in p})
            j["description"] = j["description"] or p.get("description", "")  # Bundesagentur text comes from details
        meta["assessed"] = [i for i in prev["meta"].get("assessed", []) if i in ids]
    return {"meta": meta, "jobs": {j["id"]: j for j in kept}}, skipped_applied


def cmd_pool(paths: list, out: Path, top: int):
    raw = load_items(paths)
    prev = json.loads((out / "pool.json").read_text(encoding="utf-8")) if (out / "pool.json").exists() else None
    pool, skipped_applied = make_pool(raw, dt.date.today(), prev)
    meta, kept = pool["meta"], list(pool["jobs"].values())
    out.mkdir(parents=True, exist_ok=True)
    (out / "pool.json").write_text(json.dumps(pool, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: v for k, v in meta.items() if k != "assessed"}, ensure_ascii=False))
    # a run stops at maxItems, and a job is credited to the first search that found it, so list the searches
    # that added nothing: they are candidates for a second, smaller run
    planned = json.loads((SKILL / "search_input.json").read_text(encoding="utf-8"))["keyword"]
    seen = {str(x.get("searchString") or "").rsplit(" - ", 1)[0] for x in raw}
    silent = [k for k in planned if k not in seen] if any(not x.get("source") for x in raw) else []
    if silent:
        print("Searches that added no jobs: " + json.dumps(silent, ensure_ascii=False))
    if skipped_applied:
        print("Already applied (skipped): " + "; ".join(skipped_applied))
    print("id (ba-/fh-/an- = Bundesagentur, freehire, Arbeitnow) | pre | type | title | company ([applied here] = he "
          "applied to this company before) | location | age | level | yrs | kw%")
    shown = kept[:top]
    for name, _ in CATEGORIES + [("Full-time", "")]:  # the best of each type too, so rarer types aren't buried
        shown += [j for j in kept if j["category"] == name and j not in shown][:12]
    for src in SOURCES.values():  # and of each free source: without a description at first, they pre-score lower
        shown += [j for j in kept if j["source"] == src and j not in shown][:12]
    for j in shown:
        mark = " [applied here]" if j["applied_company"] else ""
        print(f"{j['id']} | {j['prescore']} | {j['category'][:5]} | {j['title'][:75]} | {j['company'][:28]}{mark} | "
              f"{place(j['locations'][0])[:28]}{' +' if len(j['locations']) > 1 else ''} | {j['age']}d | "
              f"{j['level'][:11]} | {j['min_years'] if j['min_years'] is not None else '-'} | {j['kw']}")


def fetch(url: str, timeout: int = 15) -> tuple:
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.status, r.read(600_000).decode("utf-8", "replace")


def parse_posting(page: str) -> tuple:
    """(closed?, company apply URL or None) from LinkedIn's public job-posting HTML."""
    closed = bool(re.search(r"No longer accepting applications|closed-job|nicht mehr angenommen", page, re.I))
    m = re.search(r'<code id="applyUrl"[^>]*>\s*<!--"(.*?)"-->', page, re.S)
    if not m:
        return closed, None
    url = html.unescape(m.group(1))
    return closed, urllib.parse.parse_qs(urllib.parse.urlparse(url).query).get("url", [url])[0]


def linkedin_status(job_id: str) -> tuple:
    for attempt in range(3):
        try:
            code, page = fetch(f"https://www.linkedin.com/jobs-guest/jobs/api/jobPosting/{job_id}")
        except urllib.error.HTTPError as e:
            code, page = e.code, ""
        except (urllib.error.URLError, TimeoutError, ConnectionError):
            code, page = 0, ""
        if code != 429:
            break
        time.sleep(10 * (attempt + 1))  # LinkedIn rate limit: back off, then try again
    if code in (404, 410):
        return "Closed", None
    if code != 200:
        return "Unconfirmed", None
    closed, apply_url = parse_posting(page)
    return ("Closed" if closed else "Verified open"), apply_url


def apply_check(url: str) -> str:
    try:
        code, page = fetch(url, timeout=12)
    except urllib.error.HTTPError as e:
        return f"HTTP {e.code}"
    except Exception:
        return "unreachable"
    return "page says closed" if CLOSED_PAGE.search(page) else f"HTTP {code}"


def http_json(url: str, headers: dict = None, timeout: int = 25) -> tuple:
    """(status, JSON or None). Python can't verify rest.arbeitsagentur.de's certificate chain, so a verification
    failure retries with curl, which still checks the certificate."""
    try:
        req = urllib.request.Request(url, headers={**HEADERS, "Accept": "application/json", **(headers or {})})
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, None
    except urllib.error.URLError as e:
        if not (CURL and isinstance(e.reason, ssl.SSLCertVerificationError)):
            return 0, None
    except (TimeoutError, ConnectionError, ValueError):
        return 0, None
    args = [CURL, "-s", "-L", "--max-time", str(timeout), "-w", "\n%{http_code}"]
    for k, v in (headers or {}).items():
        args += ["-H", f"{k}: {v}"]
    try:
        # no console window per request when the desktop app (pythonw, which has no console) calls this
        out = subprocess.run(args + [url], capture_output=True, timeout=timeout + 5,
                             creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0)).stdout.decode("utf-8", "replace")
        body, _, code = out.rpartition("\n")
        return int(code), (json.loads(body) if code.startswith("2") else None)
    except (OSError, ValueError, subprocess.TimeoutExpired):
        return 0, None


def strip_html(s: str) -> str:
    s = re.sub(r"(?i)<br\s*/?>|</(?:p|li|h\d|div|ul|ol)>", "\n", s or "")
    lines = html.unescape(re.sub(r"<[^>]+>", " ", s)).splitlines()
    return "\n".join(re.sub(r"\s+", " ", l).strip() for l in lines if l.strip())


def excluded(title: str, words: list) -> bool:
    """The actor's jobTitleExclude rule (whole words, any case), applied to the free sources too."""
    return any(re.search(rf"(?<!\w){re.escape(w.lower())}(?!\w)", title.lower()) for w in words)


def kw_score(text: str, keywords: list) -> int:
    """Share of his resume keywords found (with aliases), like the actor's keywordMatchScorePercentage."""
    t = text.lower()
    hits = sum(any(re.search(rf"(?<!\w){re.escape(a.lower())}(?!\w)", t) for a in [k["keyword"], *k.get("aliases", [])])
               for k in keywords)
    return round(100 * hits / len(keywords)) if keywords else 0


def fetch_ba(keywords: list, days: int) -> list:
    """The Bundesagentur für Arbeit's job search, Germany's largest job database. Search results carry no
    description; details adds it for the jobs that get a full read."""
    found = {}
    for kw in keywords:
        for page in (1, 2, 3):
            q = urllib.parse.urlencode({"was": kw, "veroeffentlichtseit": days, "size": 100, "page": page,
                                        "zeitarbeit": "false"})
            rows = (http_json(f"{BA_API}/pc/v6/jobs?{q}", BA_KEY)[1] or {}).get("ergebnisliste") or []
            for x in rows:
                ref = x.get("referenznummer") or x.get("refnr")
                if not ref or ref in found or x.get("stellenangebotsart") in ("AUSBILDUNG", "SELBSTAENDIGKEIT"):
                    continue
                adr = (x.get("stellenlokationen") or [{}])[0].get("adresse") or {}
                part_time = any(v is True for k, v in x.items() if k.startswith("arbeitszeitTeilzeit"))
                found[ref] = {
                    "jobId": f"ba-{ref}", "source": SOURCES["ba"], "ref": ref,
                    "jobTitle": x.get("stellenangebotsTitel") or x.get("titel") or x.get("hauptberuf") or "",
                    "companyName": x.get("firma") or x.get("arbeitgeber") or "",
                    "location": german_location(adr.get("ort", ""), adr.get("region", "")),
                    "publishedAt": x.get("datumErsteVeroeffentlichung") or "",
                    "experienceLevel": "Internship" if x.get("stellenangebotsart") == "PRAKTIKUM_TRAINEE" else "",
                    "contractType": "Full-time" if x.get("arbeitszeitVollzeit") else "Part-time" if part_time else "",
                    "applyType": "EXTERNAL", "searchString": f"Bundesagentur: {kw}",
                    "dynamicFilterMatch": False if x.get("istPrivateArbeitsvermittlung") else None,
                    "jobUrl": f"https://www.arbeitsagentur.de/jobsuche/jobdetail/{ref}", "jobDescription": ""}
            if len(rows) < 100:
                break
            time.sleep(0.3)
    return list(found.values())


def fetch_freehire(days: int, pages: int = 10) -> list:
    """freehire.me collects jobs from company career sites (JOIN, Personio, Greenhouse and more). Its keyless
    API is meant for personal, non-commercial use, which this is."""
    found = {}
    for q in FREEHIRE_QUERIES:
        for page in range(pages):
            params = urllib.parse.urlencode({"q": q, "countries": "DE", "posted_within_days": days, "limit": 100,
                                             "offset": 100 * page, "description_format": "text",
                                             "sort": "posted_at", "order": "desc"})
            code, d = http_json(f"{FREEHIRE}/agent/jobs/search?{params}")
            if code == 429:  # rate limit: wait once, then retry this page
                time.sleep(30)
                code, d = http_json(f"{FREEHIRE}/agent/jobs/search?{params}")
            rows = (d or {}).get("data") or []
            for x in rows:
                e, title = x.get("enrichment") or {}, x.get("title") or ""
                if (x.get("closed_at") or e.get("seniority") in SENIOR_LEVELS
                        or (x.get("reality") or {}).get("class") in ("stale", "likely-evergreen")
                        or (x.get("is_tech") == "non_tech" and not TECH_TITLE.search(title.lower()))):
                    continue
                slug, city = x.get("public_slug"), (x.get("cities") or [""])[0]
                desc = x.get("description") or ""
                if "remote" in (x.get("location") or "").lower():
                    desc = "Fully remote (freehire lists the location as remote).\n" + desc
                found.setdefault(slug, {
                    "jobId": f"fh-{slug}", "source": SOURCES["freehire"], "slug": slug, "jobTitle": title,
                    "companyName": x.get("company") or "", "location": german_location(city) if city else "Germany",
                    "publishedAt": x.get("posted_at") or "",
                    "experienceLevel": {"intern": "Internship", "junior": "Entry level",
                                        "middle": "Associate"}.get(e.get("seniority"), ""),
                    "contractType": e.get("employment_type") or "", "applyType": "EXTERNAL",
                    "searchString": f"freehire: {q}", "jobUrl": x.get("url") or "",
                    # freehire also collects from job boards; only a career-site link counts as the company's
                    "company_url": None if JOB_BOARDS.search(x.get("url") or "") else x.get("url"),
                    "jobDescription": desc})
            if len(rows) < 100:
                break
            time.sleep(0.5)
    return list(found.values())


def fetch_arbeitnow(days: int, pages: int = 15) -> list:
    """Arbeitnow's free job feed (German jobs from company career sites), newest first, 300 per page."""
    cutoff, found = time.time() - days * 86400, []
    for page in range(1, pages + 1):
        d = http_json(f"{ARBEITNOW}?page={page}")[1] or {}
        for x in d.get("data") or []:
            if (x.get("created_at") or 0) < cutoff:
                return found
            loc = x.get("location") or ""
            parts = [p.strip() for p in loc.split(",") if p.strip() and p.strip().lower() not in ("germany", "deutschland")]
            if not (re.search(r"germany|deutschland", loc, re.I) or any(english_state(p) for p in parts)):
                continue  # the feed also carries jobs elsewhere in Europe
            types = " ".join(x.get("job_types") or []).lower()
            desc = strip_html(x.get("description") or "")
            if x.get("remote"):
                desc = "Fully remote (Arbeitnow lists this job as remote).\n" + desc
            found.append({
                "jobId": f"an-{x['slug']}", "source": SOURCES["arbeitnow"], "jobTitle": x.get("title") or "",
                "companyName": x.get("company_name") or "",
                "location": german_location(parts[0] if parts else "", parts[1] if len(parts) > 1 else ""),
                "publishedAt": dt.datetime.fromtimestamp(x["created_at"]).astimezone().isoformat(),
                "experienceLevel": "Internship" if re.search(r"intern|praktik", types) else
                "Mid-Senior level" if re.search(r"senior|lead", types) else "Associate" if re.search(r"\bmid\b", types)
                else "Entry level" if re.search(r"entry|junior|graduate", types) else "",
                "contractType": (x.get("job_types") or [""])[-1], "applyType": "EXTERNAL",
                "searchString": "Arbeitnow: newest jobs", "jobUrl": x.get("url") or "", "jobDescription": desc})
        if not (d.get("links") or {}).get("next"):
            break
        time.sleep(1.5)  # the API allows 50 requests a minute
    return found


def apify(path: str, token: str, body: dict = None, timeout: int = 60):
    """One Apify API call; the token goes in a header, never in the URL."""
    req = urllib.request.Request(f"{APIFY}{path}", method="POST" if body is not None else "GET",
                                 data=json.dumps(body).encode() if body is not None else None,
                                 headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))


def apify_run(actor: str, payload: dict, token: str, max_usd: float, progress=None) -> list:
    """Runs an actor to the end and returns its dataset. maxTotalChargeUsd caps what it may bill."""
    q = urllib.parse.urlencode({"maxTotalChargeUsd": max_usd, "timeout": 3600})
    run = apify(f"/acts/{actor.replace('/', '~')}/runs?{q}", token, payload)["data"]
    while run["status"] in ("READY", "RUNNING", "TIMING-OUT", "ABORTING"):
        time.sleep(10)
        run = apify(f"/actor-runs/{run['id']}", token)["data"]
        if progress:
            progress(run)
    items, offset = [], 0
    while run.get("defaultDatasetId"):
        q = urllib.parse.urlencode({"format": "json", "clean": 1, "offset": offset, "limit": 1000})
        page = apify(f"/datasets/{run['defaultDatasetId']}/items?{q}", token, timeout=180)
        items += page
        if len(page) < 1000:
            break
        offset += 1000
    return items


def apify_credit(token: str) -> tuple:
    """(USD used this month, monthly limit) for the Apify account."""
    d = apify("/users/me/limits", token)["data"]
    return d["current"]["monthlyUsageUsd"], d["limits"]["maxMonthlyUsageUsd"]


def fetch_linkedin(cfg: dict, token: str, days: int, max_items: int, progress=None) -> list:
    """The LinkedIn search from search_input.json, as search URLs so any window (in days) works."""
    window = f"r{days * 86400}"
    urls = [{"url": "https://www.linkedin.com/jobs/search/?" + urllib.parse.urlencode(
        {"keywords": k, "location": "Germany", "f_TPR": window})} for k in cfg["keyword"]]
    urls += [{"url": re.sub(r"f_TPR=r\d+", f"f_TPR={window}", u["url"])} for u in cfg["startUrls"]]
    payload = {k: v for k, v in cfg.items() if k not in ("keyword", "locations", "startUrls", "publishedAt")}
    payload.update(startUrls=urls, maxItems=max_items)
    items = apify_run(ACTORS["linkedin"], payload, token, round(max_items * PRICE["linkedin"] + 0.05, 2), progress)
    now = dt.datetime.now().astimezone()
    for x in items:  # LinkedIn gives only a date, plus "3 hours ago" for fresh jobs: keep that time
        m = re.search(r"(\d+)\s*(minute|hour)", x.get("postedTime") or "")
        if m:
            x["publishedAt"] = (now - dt.timedelta(**{m.group(2) + "s": int(m.group(1))})).isoformat(timespec="minutes")
    fields = LINKEDIN_FIELDS.split(",")
    return [{k: x.get(k) for k in fields} for x in items]


def fetch_board(name: str, token: str, days: int, per_query: int = 40) -> list:
    """Indeed or StepStone: one actor run per search in BOARD_QUERIES."""
    found = {}
    for q in BOARD_QUERIES:
        if name == "indeed":
            window = next((d for d in ("1", "3", "7", "14") if int(d) >= days), "14")
            payload = {"country": "de", "title": q, "location": "Deutschland", "limit": per_query, "datePosted": window}
        else:
            payload = {"keywords": q, "ag": "age_1" if days <= 1 else "age_7", "limit": per_query}
        for x in apify_run(ACTORS[name], payload, token, round(per_query * PRICE[name] + 0.02, 2)):
            item = indeed_item(x, q) if name == "indeed" else stepstone_item(x, q)
            if item:
                found.setdefault(item["jobId"], item)
    return list(found.values())


def indeed_item(x: dict, q: str):
    if x.get("expired") or not x.get("key"):
        return None
    loc, pay = x.get("location") or {}, x.get("baseSalary") or {}
    city = loc.get("city") or ""
    salary = (f"{pay.get('min') or ''}–{pay.get('max') or ''} {pay.get('currencyCode') or ''} per "
              f"{(pay.get('unitOfWork') or '').lower()}") if pay.get("min") or pay.get("max") else ""
    return {"jobId": f"in-{x['key']}", "source": SOURCES["indeed"], "jobTitle": x.get("title") or "",
            "companyName": (x.get("employer") or {}).get("name") or "",
            "location": ", ".join(p for p in (city, STATE_CODES.get(loc.get("admin1Code") or "")
                                              or CITY_STATES.get(city.lower(), ""), "Germany") if p),
            "publishedAt": x.get("datePublished") or x.get("dateOnIndeed") or "", "applyType": "EXTERNAL",
            "searchString": f"Indeed: {q}", "jobUrl": x.get("url") or "",
            "company_url": None if JOB_BOARDS.search(x.get("jobUrl") or "") else x.get("jobUrl"),
            "salaryInfo": [salary] if salary else [], "jobDescription": (x.get("description") or {}).get("text") or ""}


def stepstone_item(x: dict, q: str):
    if not x.get("id"):
        return None
    sections = x.get("textSections") or []
    desc = "\n".join(f"{s.get('title') or ''}:\n{strip_html(s.get('content') or '')}" for s in sections)
    links = [m for s in sections for m in re.findall(r'href="(https?://[^"]+)"', s.get("content") or "")]
    if str(x.get("workFromHome") or "0") != "0":
        desc = "Home office possible (StepStone).\n" + desc
    pay = x.get("salary") or {}
    salary = (f"{pay.get('min') or ''}–{pay.get('max') or ''} {pay.get('currency') or ''} per {pay.get('period') or ''}"
              if (pay.get("min") or pay.get("max")) and not pay.get("isPredicted") else "")
    return {"jobId": f"ss-{x['id']}", "source": SOURCES["stepstone"], "jobTitle": x.get("title") or "",
            "companyName": (x.get("company") or {}).get("name") or "",
            "location": german_location(((x.get("location") or {}).get("location") or "").split(",")[0]),
            "publishedAt": x.get("datePosted") or "", "applyType": "EXTERNAL",
            "searchString": f"StepStone: {q}", "jobUrl": x.get("url") or "",
            "company_url": next((u for u in links if not JOB_BOARDS.search(u) and "stepstone." not in u), None),
            "salaryInfo": [salary] if salary else [], "jobDescription": desc or x.get("textSnippet") or ""}


def keep_tech(x: dict, cfg: dict) -> bool:
    """Scores a free-source job against his resume keywords; keeps tech jobs whose titles aren't excluded."""
    x["keywordMatchScorePercentage"] = kw_score(x["jobTitle"] + "\n" + x["jobDescription"], cfg["resumeKeywords"])
    return not excluded(x["jobTitle"], cfg["jobTitleExclude"]) and bool(
        TECH_TITLE.search(x["jobTitle"].lower()) or x["keywordMatchScorePercentage"] >= 40)


def fetch_sources(names: list, days: int, token: str = "", linkedin_max: int = 400, progress=print) -> tuple:
    """(items, counts) from the named sources, searched in parallel. LinkedIn, Indeed and StepStone run on
    Apify and need a token; LinkedIn's actor filters itself, the others go through keep_tech."""
    cfg = json.loads((SKILL / "search_input.json").read_text(encoding="utf-8"))
    run = {"ba": lambda: fetch_ba(cfg["keyword"], days), "freehire": lambda: fetch_freehire(days),
           "arbeitnow": lambda: fetch_arbeitnow(days), "indeed": lambda: fetch_board("indeed", token, days),
           "stepstone": lambda: fetch_board("stepstone", token, days),
           "linkedin": lambda: fetch_linkedin(cfg, token, days, linkedin_max)}

    def one(name):
        if name in ACTORS and not token:
            return name, [], "skipped (no Apify token)"
        try:
            got = run[name]()
        except Exception as e:  # one failing source must not sink the whole search
            return name, [], f"failed: {e}"
        keep = got if name == "linkedin" else [x for x in got if keep_tech(x, cfg)]
        progress(f"{SOURCES[name]}: {len(keep)} jobs")
        return name, keep, f"{len(keep)} tech jobs kept of {len(got)} found"

    items, counts = [], {}
    with ThreadPoolExecutor(len(names) or 1) as ex:
        for name, keep, note in ex.map(one, names):
            items += keep
            counts[SOURCES[name]] = note
    return items, counts


def cmd_fetch(out: Path, days: int, sources: list):
    items, counts = fetch_sources(sources, days, os.environ.get("APIFY_TOKEN", ""), progress=lambda m: None)
    out.mkdir(parents=True, exist_ok=True)
    (out / "free_sources.json").write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(counts, ensure_ascii=False))
    print(f"Wrote {out / 'free_sources.json'}; pass it to pool together with the Apify dump(s)")


def ba_details(ref: str) -> tuple:
    """(status, description) from the Bundesagentur's job details; 404 means the posting was taken down."""
    code, d = http_json(f"{BA_API}/pc/v4/jobdetails/{base64.b64encode(ref.encode()).decode()}", BA_KEY)
    if code == 404:
        return "Closed", ""
    if code != 200 or not d:
        return "Unconfirmed", ""
    return "Verified open", d.get("stellenangebotsBeschreibung") or d.get("stellenbeschreibung") or ""


def freehire_status(slug: str) -> str:
    code, d = http_json(f"{FREEHIRE}/jobs/{slug}")
    if code == 404 or ((d or {}).get("data") or {}).get("closed_at"):
        return "Closed"
    return "Verified open" if code == 200 else "Unconfirmed"


def page_status(url: str) -> str:
    check = apply_check(url)
    return ("Closed" if check in ("HTTP 404", "HTTP 410", "page says closed")
            else "Verified open" if check == "HTTP 200" else "Unconfirmed")


def workplace(desc: str) -> str:
    d = desc.lower()
    remote = re.search(r"remote|home ?office|home-office|mobiles? arbeiten|mobile work|from home|von zu hause", d)
    if re.search(r"(?:100\s?%|fully|vollständig|komplett|full) ?remote|remote[- ]first", d):
        return "Remote"
    if "hybrid" in d or (remote and re.search(r"vor ort|on-?site|präsenz|im büro|in the office|in our office", d)):
        return "Hybrid"
    if remote:
        return "Hybrid (remote possible)"
    return "On-site" if re.search(r"vor ort|on-?site|präsenz", d) else "Not stated"


def sections(desc: str) -> dict:
    cur, out = None, {"tasks": [], "needs": []}
    for line in (l.strip(" \t•·-–*●▪✓✔") for l in desc.splitlines()):
        if not line:
            continue
        heading = len(line) <= 60 or (len(line) <= 120 and line.endswith(":"))
        if heading and OTHER_HEAD.match(line):  # checked first: "your working environment" isn't a task list
            cur = None
        elif heading and TASK_HEAD.match(line):
            cur = "tasks"
        elif heading and REQ_HEAD.match(line):
            cur = "needs"
        elif cur:
            out[cur].append(line)
    return out


def snippets(j: dict) -> dict:
    desc = j["description"]
    # by line, and by sentence too, because some postings (a fifth of freehire's) come as one paragraph
    lines = [l.strip(" \t•·-–*●▪✓✔") for l in re.split(r"\n|(?<=[.!?])\s+(?=[A-ZÄÖÜ])", desc) if l.strip()]

    def grab(rx, n, extra=None):
        hits = [l[:170] for l in lines if re.search(rx, l, re.I) and (not extra or re.search(extra, l, re.I))]
        return list(dict.fromkeys(hits))[:n]

    sec = sections(desc)
    return {"tasks": " • ".join(sec["tasks"])[:300], "needs": " • ".join(sec["needs"])[:600],
            "language": grab(r"deutsch|german|englisch|english", 3,
                             r"fließend|fliessend|verhandlungssicher|sehr gut|gut|grundkenntnisse|muttersprach|"
                             r"\b[abc][12]\b|fluent|proficien|business|native|good|excellent|level|kenntnisse|skills|"
                             r"spoken|written|wort und schrift|required|erforderlich|vorteil|plus"),
            "experience": grab(r"erfahrung|experience", 2, r"\d|erste|first|mehrjährig|several|proven|fundiert|einschlägig"),
            "hours": grab(r"\b\d{1,2}\s*(?:-|–|bis|to)?\s*\d{0,2}\s*(?:std\.?|stunden|hours|h)\b", 2),
            "start": grab(r"ab sofort|asap|as soon as possible|eintritt|beginn|starttermin|start date|befristet|"
                          r"\d+\s*(?:monate|months)", 2),
            "salary": [s for s in j["salary_info"] if s][:1] + grab(
                r"gehalt|vergütung|salary|stundenlohn|brutto|compensation|remuneration|"
                r"(?:€|eur|euro)\s*(?:pro|per|je|/)|(?:pro|per|je)\s*(?:stunde|hour|monat|month)", 2),
            "enrolment": grab(r"immatrikul|eingeschrieben|enrolled|enrolment|enrollment", 1)}


def check(jobs: dict, ids: list, progress=None):
    """Checks each job on its own source (open or closed, company apply link) and adds the snippets that
    scoring needs. Updates the job dicts in place."""
    for n, i in enumerate(ids, 1):
        j, src = jobs[i], jobs[i].get("source", "LinkedIn")
        if progress:
            progress(n, len(ids))
        if src == "LinkedIn":
            j["status"], j["apply_url"] = linkedin_status(i)
            time.sleep(1.2)  # stay well below LinkedIn's rate limit
        elif src == SOURCES["ba"]:
            j["status"], desc = ba_details(j["ref"])
            j["description"], j["apply_url"] = j["description"] or desc, None
        else:
            j["status"] = freehire_status(j["slug"]) if src == SOURCES["freehire"] else page_status(j["url"])
            if j["status"] == "Unconfirmed" and src in (SOURCES["indeed"], SOURCES["stepstone"]):
                j["status"] = "Listed today"  # their pages block scripts, but the job was in today's live search
            j["apply_url"] = None
        # a company link from another source's copy of the job (see dedupe) beats LinkedIn's apply button
        j["apply_url"] = j["apply_url"] or j.get("company_url")
        j["lang"], j["domain"] = posting_language(j["description"]), domain(j["title"], j["description"])
        j["checked"] = dt.datetime.now().isoformat(timespec="minutes")
    with ThreadPoolExecutor(8) as ex:  # company sites are different hosts, so these can run in parallel
        checks = list(ex.map(lambda i: apply_check(jobs[i]["apply_url"]) if jobs[i]["apply_url"] else "", ids))
    for i, result in zip(ids, checks):
        j = jobs[i]
        j["apply_check"] = result
        if result == "page says closed":
            j["status"] = "Closed"
        j["snippets"], j["workplace_hint"] = snippets(j), workplace(j["description"])


def describe(i: str, j: dict) -> str:
    s, src = j["snippets"], j.get("source", "LinkedIn")
    apply = (f"company site, {j['apply_check']}" if j["apply_url"] else
             "Easy Apply" if j["apply_type"] == "EASY_APPLY" else f"via {src}")
    mark = " [he applied to this company before]" if j.get("applied_company") else ""
    also = f" (also on {', '.join(j['also_on'])})" if j.get("also_on") else ""
    lines = [f"\n[{i}] {j['title']} | {j['company']}{mark} | {' / '.join(place(l) for l in j['locations'])} | "
             f"posted {j['published'] or '?'} ({j['age']}d) | {j['category']} | {j['level']}, {j['contract']} | "
             f"{j['lang']} posting | applicants: {j['applicants'] or '?'} | source: {src}{also}",
             f"  status: {j['status']} | apply: {apply} | workplace: {j['workplace_hint']}"]
    lines += [f"  {k}: {s[k]}" for k in ("tasks", "needs") if s[k]]
    lines += [f"  {k}: " + " || ".join(s[k]) for k in ("language", "experience", "hours", "start", "salary", "enrolment")
              if s[k]]
    return "\n".join(lines)


def cmd_details(ids: list, out: Path):
    pool = json.loads((out / "pool.json").read_text(encoding="utf-8"))
    jobs = pool["jobs"]
    unknown = [i for i in ids if i not in jobs]
    if unknown:
        print("Not in pool.json (dropped or mistyped): " + ", ".join(unknown))
    ids = [i for i in ids if i in jobs]
    check(jobs, ids)
    for i in ids:
        print(describe(i, jobs[i]))
    pool["meta"]["assessed"] = sorted(set(pool["meta"].get("assessed", [])) | set(ids))
    (out / "pool.json").write_text(json.dumps(pool, ensure_ascii=False, indent=1), encoding="utf-8")


COLUMNS = [("Rank", 6, "rank"), ("Match score", 8, "score"), ("Job title", 42, "title"), ("Company", 26, "company"),
           ("Location", 24, "location"), ("Workplace type", 13, "workplace"), ("Employment type", 13, "employment_type"),
           ("Technical domain", 15, "domain"), ("Posted date", 11, "posted"), ("Experience required", 32, "experience"),
           ("Language requirements", 32, "language"), ("Salary", 18, "salary"), ("Why it matches", 55, "why"),
           ("Missing skills", 36, "missing"), ("Eligibility", 30, "eligibility"), ("Vacancy status", 14, "status"),
           ("LinkedIn link", 15, "linkedin"), ("Apply link", 22, "apply")]
SHEETS = [("Top 30 Jobs", None), ("Working Student & Internship", {"Werkstudent", "Internship"}),
          ("Master's Thesis & Research", {"Thesis", "Research"}), ("Junior & Full-Time", {"Full-time", "Trainee"})]


def write_xlsx(rows: list, summary: list, path: Path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill

    body, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    link = Font(name="Arial", size=10, color="0563C1", underline="single")
    fill, wrap = PatternFill("solid", fgColor="1F4E78"), Alignment(wrap_text=True, vertical="top")
    wb = Workbook()
    wb.remove(wb.active)
    for name, types in SHEETS:
        ws = wb.create_sheet(name)
        ws.append([c[0] for c in COLUMNS])
        for r in (r for r in rows if not types or r["employment_type"] in types):
            values = [r[key] for _, _, key in COLUMNS]
            values[COLUMNS.index(("Posted date", 11, "posted"))] = (dt.date.fromisoformat(r["posted"])
                                                                    if r["posted"] else "Unavailable")
            ws.append(values)
            n = ws.max_row
            if r["linkedin"]:
                ws.cell(n, 17).value, ws.cell(n, 17).hyperlink = "Open on LinkedIn", r["linkedin"]
            else:
                ws.cell(n, 17).value = f"Not on LinkedIn (from {r['source']})"
            ws.cell(n, 18).value, ws.cell(n, 18).hyperlink = r["apply_label"], r["apply"]
        for row in ws.iter_rows():
            for cell in row:
                cell.font, cell.alignment = (bold if cell.row == 1 else link if cell.hyperlink else body), wrap
                if cell.row == 1:
                    cell.fill = fill
                if cell.column == 9 and cell.row > 1 and isinstance(cell.value, dt.date):
                    cell.number_format = "yyyy-mm-dd"
        for i, (_, width, _) in enumerate(COLUMNS, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        ws.freeze_panes, ws.auto_filter.ref = "C2", ws.dimensions
        if ws.max_row > 1:
            ws.conditional_formatting.add(f"B2:B{ws.max_row}", ColorScaleRule(
                start_type="num", start_value=50, start_color="F8696B", mid_type="num", mid_value=70,
                mid_color="FFEB84", end_type="num", end_value=90, end_color="63BE7B"))
    ws = wb.create_sheet("Search Summary")
    ws.append(["Item", "Value", "Notes"])
    for item in summary:
        ws.append(item)
    for row in ws.iter_rows():
        for cell in row:
            cell.font, cell.alignment = (bold if cell.row == 1 else body), wrap
            if cell.row == 1:
                cell.fill = fill
    for col, width in zip("ABC", (34, 80, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    wb.save(path)


def cmd_export(out: Path):
    pool = json.loads((out / "pool.json").read_text(encoding="utf-8"))
    picks = json.loads((out / "picks.json").read_text(encoding="utf-8"))
    meta, rows = pool["meta"], []
    for p in picks["picks"]:
        j = pool["jobs"][p["id"]]
        if j.get("status") == "Closed":
            print(f"Skipped, closed: {j['title']} | {j['company']}")
            continue
        scores = {k: max(0, min(w, p["scores"].get(k, 0))) for k, w in WEIGHTS.items()}
        src, company = j.get("source", "LinkedIn"), j.get("apply_url") or j.get("company_url")
        apply = company or j.get("url") or j["linkedin"]
        label = ("Apply on company site" if company else "Easy Apply on LinkedIn"
                 if src == "LinkedIn" and j["apply_type"] == "EASY_APPLY" else f"Apply via {src}")
        rows.append({"id": j["id"], "source": src, "score": sum(scores.values()), "scores": scores, "title": j["title"],
                     "company": j["company"], "location": " / ".join(place(l) for l in j["locations"]),
                     "states": [state(l) or "Germany (nationwide)" for l in j["locations"]],
                     "workplace": p.get("workplace") or j.get("workplace_hint", "Not stated"),
                     "employment_type": p.get("employment_type") or j["category"],
                     "domain": p.get("domain") or j["domain"], "posted": j["published"], "age": j["age"],
                     "experience": p.get("experience", ""), "language": p.get("language", ""),
                     "salary": p.get("salary") or "Not stated", "why": p.get("why", ""),
                     "missing": p.get("missing", ""), "eligibility": p.get("eligibility", "Needs verification"),
                     "status": j.get("status", "Unconfirmed"), "linkedin": j["linkedin"], "apply": apply,
                     "apply_label": label, "posting_language": j["lang"], "description": j["description"]})
    rows.sort(key=lambda r: (-r["score"], r["age"]))
    rows = rows[:30]
    for n, r in enumerate(rows, 1):
        r["rank"] = n
    date, s = meta["date"], picks.get("summary", {})
    assessed = [pool["jobs"][i] for i in meta["assessed"] if i in pool["jobs"]]
    count = lambda key: ", ".join(f"{k} {v}" for k, v in Counter(r[key] for r in rows).most_common())
    by_source = lambda key: ", ".join(f"{k} {v}" for k, v in meta.get(key, {}).items())
    summary = [
        ["Search date", date, "Posting dates as each source states them; ages count from this date"],
        ["Search tools used", s.get("tools", "LinkedIn through Apify (cheap_scraper/linkedin-job-scraper), "
                                             "Bundesagentur für Arbeit job search API, freehire API, Arbeitnow API"),
         "Vacancy status checked on each job's source and on the company apply page"],
        ["Number of searches performed", s.get("searches", meta["searches_with_results"]),
         f"{meta['searches_with_results']} LinkedIn searches returned results"],
        ["Total postings retrieved", meta["retrieved"], s.get("window", "Posted in the last 7 days")],
        ["Postings per source", by_source("retrieved_by_source") or "LinkedIn " + str(meta["retrieved"]),
         "After filters: " + (by_source("kept_by_source") or str(meta["kept"]))],
        ["Unique postings after deduplication", meta["unique"], "Same ID, or same company and title on any source"],
        ["Removed by filters", ", ".join(f"{k}: {v}" for k, v in meta["dropped"].items()) or "none",
         "Language requirements never removed a job"],
        ["Number assessed", len(assessed), "Read in full and checked on their source"],
        ["Number verified open", sum(j.get("status") == "Verified open" for j in assessed),
         "Open on its source at the check time; company pages that say a job is closed count as closed"],
        ["Number shortlisted", len(rows), "Ranked by the weighted match score (language is not scored)"],
        ["Domain coverage", count("domain"), "Pool: " + ", ".join(f"{k} {v}" for k, v in meta["domains"].items())],
        ["Employment type coverage", count("employment_type"),
         "Pool: " + ", ".join(f"{k} {v}" for k, v in meta["categories"].items())],
        ["Geographic coverage", "Shortlist: " + ", ".join(f"{k} {v}" for k, v in Counter(
            s for r in rows for s in set(r["states"])).most_common()),
         "Pool: " + ", ".join(f"{k} {v}" for k, v in sorted(meta["states"].items(), key=lambda kv: -kv[1]))],
        ["German- and English-language postings", count("posting_language"),
         "Pool: " + ", ".join(f"{k} {v}" for k, v in meta["languages"].items())],
        ["Known search limitations", s.get("limitations", ""), ""],
        ["Apify cost", s.get("cost", ""), "Pay per result, free tier; the other sources are free"],
        ["Build a CV and cover letter", "In Claude Code: /job-apply <Rank> (for example /job-apply 3)",
         f"Uses {out.as_posix()}/jobs.json, which holds each full posting"],
    ]
    out.mkdir(parents=True, exist_ok=True)
    (out / "jobs.json").write_text(json.dumps({"date": date, "jobs": rows}, ensure_ascii=False, indent=1),
                                   encoding="utf-8")
    write_xlsx(rows, summary, out / f"Veer_Job_Search_{date}.xlsx")
    board = [{k: r[k] for k in ("rank", "score", "title", "company", "location", "employment_type", "workplace",
                                "domain", "posted", "status", "why", "missing", "experience", "language", "salary",
                                "eligibility", "linkedin", "apply", "source")} for r in rows]
    rel = (out.resolve().relative_to(PROJECT) if out.resolve().is_relative_to(PROJECT) else out).as_posix()
    tpl = (SKILL / "assets" / "job-board.html").read_text(encoding="utf-8")
    # "</" inside the JSON would end the <script> element early
    (out / "job_board.html").write_text(tpl.replace("__JOBS__", json.dumps(board, ensure_ascii=False)
                                                    .replace("</", "<\\/")).replace("__FILE__", f"{rel}/jobs.json"),
                                        encoding="utf-8")
    print(f"Wrote {out / 'jobs.json'}, {out / f'Veer_Job_Search_{date}.xlsx'} and {out / 'job_board.html'}")
    print("| # | Score | Job | Company | Location | Type | Posted | Status |\n|---|---|---|---|---|---|---|---|")
    for r in rows:
        print(f"| {r['rank']} | {r['score']} | {r['title']} | {r['company']} | {r['location']} | "
              f"{r['employment_type']} | {r['posted']} | {r['status']} |")


def selftest():
    assert category("Werkstudent (m/w/d) Masterarbeit KI") == "Thesis"
    assert category("Studentische Hilfskraft (SHK) Informatik") == "Research"
    assert category("Working Student Data Engineering") == "Werkstudent"
    assert category("Praktikant Softwareentwicklung") == "Internship"
    assert category("Junior Data Engineer") == "Full-time"
    assert domain("Werkstudent KI & Prozessautomatisierung", "") == "AI/ML"
    assert domain("Junior Cloud Engineer", "") == "Cloud/DevOps"
    assert min_years(["2-5", "3+"]) == 2 and min_years(["100"]) is None and min_years([]) is None
    assert years_in_text("Mindestens 3 Jahre Berufserfahrung in Python") == ["3"]
    assert years_in_text("over 100-year history") == []
    assert age_from_text("3 hours ago") == 0 and age_from_text("2 weeks ago") == 14
    assert place("Berlin, Berlin, Germany") == "Berlin" and place("Munich, Bavaria, Germany") == "Munich, Bavaria"
    assert state("Munich, Bavaria, Germany") == "Bavaria" and state("Germany") == ""
    a = {"id": "1", "company": "Thales", "title": "Werkstudent (m/f/d) Data AI", "locations": ["Ulm"], "age": 3}
    b = dict(a, id="2", title="Werkstudent (w/m/d) Data AI", locations=["Kiel"], age=1)
    d = dedupe([a, b])
    assert len(d) == 1 and d[0]["id"] == "2" and d[0]["locations"] == ["Kiel", "Ulm"]
    sets = [{"deag", "ki", "automatisierung"}, {"mb", "energy", "platform", "engineer"},
            {"atlascopco", "leybold", "software", "development"}]
    deag = {"company": "DEAG Deutsche Entertainment AG", "title": "Werkstudent KI & Automatisierung (m/w/d)"}
    assert applied_match(deag, sets) == "job"
    assert applied_match({"company": "XYZ Energy GmbH", "title": "Working Student Platform Engineer"}, sets) is None
    assert applied_match({"company": "Leybold", "title": "Werkstudent Softwareentwicklung (m/w/d)"}, sets) == "company"
    assert workplace("Hybrides Arbeiten möglich") == "Hybrid" and workplace("100% remote") == "Remote"
    assert workplace("Remote möglich, Büro in Köln vor Ort") == "Hybrid" and workplace("Python") == "Not stated"
    sec = sections("Deine Aufgaben:\nBuild APIs\nDein Profil\nPython\nSQL\nWas wir bieten\nObst")
    assert sec == {"tasks": ["Build APIs"], "needs": ["Python", "SQL"]}
    sec = sections("Your working environment\nMAN is big\nYour mission as \"Werkstudent in Radar\":\nPipelines\n"
                   "We are looking forward to:\nPython\nThe Group invests €4,5 billion")
    assert sec["tasks"] == ["Pipelines"] and sec["needs"][0] == "Python"
    assert snippets({"description": "Group sales of €22.1 billion\nab 17,74 € pro Stunde", "salary_info": []}
                    )["salary"] == ["ab 17,74 € pro Stunde"]
    page = ('<code id="applyUrl" style="display: none"><!--"https://www.linkedin.com/jobs/view/externalApply/1?'
            'url=https%3A%2F%2Fjobs%2Eexample%2Ede%2F42&amp;urlHash=x"--></code>')
    assert parse_posting(page) == (False, "https://jobs.example.de/42")
    assert parse_posting('<figure class="closed-job">No longer accepting applications</figure>')[0]
    assert english_state("NORDRHEIN_WESTFALEN") == "North Rhine-Westphalia" and english_state("Thüringen") == "Thuringia"
    assert german_location("Gilching", "Bayern") == "Gilching, Bavaria, Germany" and english_state("Munich") == ""
    assert german_location("Berlin") == "Berlin, Berlin, Germany" and place(german_location("Berlin")) == "Berlin"
    assert state_label("Kiel, Germany") == "state not given" and state_label("Germany") == "nationwide/remote"
    assert strip_html("<p>Hi &amp; welcome</p><ul><li>Python</li><li>SQL</li></ul>").splitlines() == \
        ["Hi & welcome", "Python", "SQL"]
    assert excluded("Senior Data Engineer", ["Senior"]) and not excluded("Seniority Analyst", ["Senior"])
    kws = [{"keyword": "Python"}, {"keyword": "SQL", "aliases": ["PostgreSQL"]}, {"keyword": "Java"}]
    assert kw_score("Python and PostgreSQL", kws) == 67 and kw_score("JavaScript", kws) == 0
    assert posting_language("") == "Unknown"
    li = {"id": "9", "source": "LinkedIn", "company": "Thales", "title": "Werkstudent Data AI", "locations": ["Ulm"],
          "age": 3}
    fh = {"id": "fh-x", "source": "freehire", "company": "Thales GmbH", "title": "Werkstudent (m/w/d) Data AI",
          "locations": ["Ulm, Baden-Württemberg, Germany"], "age": 1, "company_url": "https://jobs.thales.example/42"}
    d = dedupe([fh, li])
    assert len(d) == 1 and d[0]["id"] == "9" and d[0]["company_url"] == fh["company_url"] and d[0]["also_on"] == ["freehire"]
    assert d[0]["locations"] == ["Ulm"]
    assert same_company(company_key("Scalable Capital"), company_key("ScalableGmbH"))
    assert same_company(company_key("nexis"), company_key("nexisiag"))
    assert not same_company(company_key("Siemens"), company_key("BASF SE"))
    assert german_location("Kaufungen, Hessen", "HESSEN") == "Kaufungen, Hesse, Germany"
    assert german_location("München") == "München, Bavaria, Germany"
    ind = indeed_item({"key": "a9", "title": "Werkstudent Data (m/w/d)", "url": "https://de.indeed.com/viewjob?jk=a9",
                       "jobUrl": "https://elunic.teamtailor.com/jobs/1", "datePublished": "2026-10-19T05:00:00.000Z",
                       "expired": False, "location": {"city": "München", "admin1Code": "BY"},
                       "employer": {"name": "elunic AG"}, "description": {"text": "Python"}}, "q")
    assert ind["location"] == "München, Bavaria, Germany" and ind["company_url"] == "https://elunic.teamtailor.com/jobs/1"
    assert indeed_item({"key": "x", "expired": True}, "q") is None
    ss = stepstone_item({"id": 7, "title": "Werkstudent Data", "url": "https://www.stepstone.de/x", "workFromHome": "2",
                         "datePosted": "2026-10-26T00:26:39+01:00", "location": {"location": "Nürnberg, Home-Office"},
                         "company": {"name": "Circana"}, "textSections": [
                             {"title": "Profil", "content": "<ul><li>Python</li></ul>"},
                             {"title": "Weitere Informationen",
                              "content": '<a href="https://circana.onlyfy.jobs/apply/1">online</a>'}]}, "q")
    assert ss["location"] == "Nürnberg, Bavaria, Germany" and ss["company_url"] == "https://circana.onlyfy.jobs/apply/1"
    sj = to_job(ss, dt.date(2026, 10, 27))
    assert sj["published"] == "2026-10-26" and sj["posted_at"] == "2026-10-26T00:26:39+01:00" and sj["age"] == 1
    assert drop_reason({**sj, "title": "Doctoral Fellowship in Data Science"}, []) == "PhD or postdoc"
    assert workplace(ss["jobDescription"]) == "Hybrid (remote possible)" and sections(ss["jobDescription"])["needs"] == ["Python"]
    assert posting_language("Wir suchen dich für unser Team und die Entwicklung") == "German"
    job = {"title": "Werkstudent Data", "kw": 60, "age": 1, "category": "Werkstudent", "level": "", "min_years": None}
    assert prescore(job) == 80
    print("selftest ok")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Merge, check and rank LinkedIn jobs for Veer.")
    ap.add_argument("command", nargs="?", choices=["fetch", "pool", "details", "export"])
    ap.add_argument("args", nargs="*", help="pool: dataset dump files; details: job IDs")
    ap.add_argument("--out", type=Path, default=PROJECT / "jobs" / dt.date.today().isoformat())
    ap.add_argument("--top", type=int, default=150, help="pool: how many candidates to print")
    ap.add_argument("--days", type=int, default=7, help="fetch: posted within this many days")
    ap.add_argument("--sources", default=",".join(FREE_SOURCES),
                    help="fetch: comma-separated, from " + ", ".join(SOURCES) + " (the last three need APIFY_TOKEN)")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        selftest()
    elif a.command == "fetch":
        cmd_fetch(a.out, a.days, [s.strip() for s in a.sources.split(",") if s.strip() in SOURCES])
    elif a.command == "pool":
        cmd_pool(a.args, a.out, a.top)
    elif a.command == "details":
        cmd_details(a.args, a.out)
    elif a.command == "export":
        cmd_export(a.out)
    else:
        ap.error("give a command: pool, details or export")
